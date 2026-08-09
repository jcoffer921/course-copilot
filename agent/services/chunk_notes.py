"""
Chunks one lecture's raw notes or slide deck into notes/<lecture_id>.json.
Two input types, same output schema (courses/<course_id>/notes/<lecture_id>.json),
tagged with "source": "notes" or "source": "slides" so downstream consumers
(ask.py, quiz.py) can tell them apart if it ever matters.

Uses AsyncAnthropic so the DRF async view can await it directly without
blocking the ASGI event loop; the management command wraps the same
coroutine with asyncio.run() for standalone CLI use — same pattern as
syllabus_extraction.py and ask.py.
"""

import io
import json
import re
from pathlib import Path

from asgiref.sync import sync_to_async

from . import storage
from .client import MODEL_DEFAULT as MODEL, get_client
from .storage import CourseNotFoundError
from .syllabus_extraction import extract_text_from_bytes

CHUNKING_SYSTEM_PROMPT = """You split one lecture's material into semantically \
meaningful study chunks — grouped by topic/subtopic within the lecture, not by \
character count or arbitrary length. A chunk should cover one coherent idea a \
student could later be quizzed on independently.

You will be given:
1. This course's official topics list, from its syllabus.
2. The lecture's raw source text (already extracted from notes or slides — \
section headings in the text, if any, are structural markers, not chunks \
themselves).

Rules:
- Use ONLY information present in the provided source text. Do not add facts, \
examples, or explanations that aren't there, and do not fill in generic \
textbook knowledge about the topic that the source itself doesn't state.
- Each chunk's "topic" should be one of the course's official topics (reuse \
that exact string) when the chunk's content clearly belongs to one of them. \
If a chunk doesn't cleanly match any listed topic, give it a short, specific \
topic name of your own instead of forcing a bad fit — but always prefer the \
official list when it genuinely applies, since topic-mastery tracking depends \
on topic names staying consistent across lectures, slide decks, and quizzes.
- Each chunk needs a short, unique, kebab-case "id" (e.g. "recursion-base-case").
- Do not create a chunk for a heading/section-divider with no real content \
under it, and do not pad a thin chunk with restated headings just to give it \
body text — if there's nothing substantive to quiz on, leave it out.
- Output ONLY valid JSON matching the schema below. No preamble, no markdown \
fences, no commentary.

Schema:
{
  "chunks": [{"id": "string", "topic": "string", "text": "string"}]
}
"""


class MalformedSourceError(ValueError):
    """Raised when the source file (notes or slide deck) has nothing usable
    to chunk — e.g. empty text, or a slide deck with no content slides."""


# --------------------------------------------------------------------------
# Source text extraction (sync — file I/O / local parsing, not worth async)
# --------------------------------------------------------------------------

def _extract_slides_from_bytes(data: bytes) -> list:
    """Returns [{"title": str, "body": str}, ...] for every slide, in order.
    Uses python-pptx's shape/text_frame API — not hand-rolled XML parsing."""
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as e:
        raise MalformedSourceError(f"couldn't read .pptx file: {e}")

    slides = []
    for slide in prs.slides:
        title = ""
        body_parts = []
        # slide.shapes.title returns a fresh wrapper object on every access, so
        # `shape is title_shape` never matches during the iteration below even
        # for the same underlying shape — compare stable shape_ids instead.
        title_shape = slide.shapes.title
        title_id = title_shape.shape_id if title_shape is not None else None
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            if shape.shape_id == title_id:
                title = text
            else:
                body_parts.append(text)
        slides.append({"title": title, "body": "\n".join(body_parts).strip()})

    if not slides:
        raise MalformedSourceError("slide deck has no slides")
    return slides


def _slides_to_source_text(slides: list) -> str:
    """Collapses divider/section-header slides (title only, no body) into a
    heading prefix carried forward onto the next content slide, instead of
    emitting them as their own near-empty chunk later."""
    parts = []
    current_section = None
    for s in slides:
        if not s["body"]:
            # Divider slide (title-only, e.g. a section header or "Questions?"
            # closer) — remember its title as context, don't emit it alone.
            current_section = s["title"] or current_section
            continue
        heading = f"{current_section} — {s['title']}" if current_section and s["title"] else (s["title"] or current_section)
        parts.append(f"## {heading}\n{s['body']}" if heading else s["body"])

    if not parts:
        raise MalformedSourceError(
            "slide deck has no slides with body content to chunk — it may be all "
            "section dividers/title slides"
        )
    return "\n\n".join(parts)


def extract_text_from_source(data: bytes, filename: str) -> tuple:
    """Returns (source_text, source_type) where source_type is 'notes' or
    'slides'. Reuses syllabus_extraction's .pdf/.txt/.md handling for the
    'notes' path rather than reimplementing it."""
    suffix = Path(filename).suffix.lower()

    if suffix == ".pptx":
        slides = _extract_slides_from_bytes(data)
        return _slides_to_source_text(slides), "slides"

    try:
        text = extract_text_from_bytes(data, filename)
    except ValueError as e:
        if "unsupported file type" in str(e):
            raise ValueError(f"{e} (or .pptx for slide decks)")
        raise
    return text, "notes"


def read_source_from_path(path: Path) -> tuple:
    """For the CLI management command — reads a local file."""
    return extract_text_from_source(path.read_bytes(), path.name)


def read_source_from_upload(upload) -> tuple:
    """For the DRF view — reads a Django UploadedFile."""
    return extract_text_from_source(upload.read(), upload.name)


# --------------------------------------------------------------------------
# Chunk id cleanup — defense against the model producing duplicate or
# malformed slugs, same spirit as storage.py's path-safety helpers
# --------------------------------------------------------------------------

def _dedupe_chunk_ids(chunks: list) -> list:
    seen = {}
    result = []
    for c in chunks:
        if not isinstance(c, dict):
            continue
        base_id = re.sub(r"[^a-z0-9]+", "-", str(c.get("id", "chunk")).lower()).strip("-") or "chunk"
        n = seen.get(base_id, 0)
        seen[base_id] = n + 1
        final_id = base_id if n == 0 else f"{base_id}-{n + 1}"
        text = str(c.get("text", "")).strip()
        if not text:
            continue  # a chunk with no real text is exactly the junk this whole flow avoids
        result.append({"id": final_id, "topic": str(c.get("topic", "")).strip(), "text": text})
    return result


# --------------------------------------------------------------------------
# Chunking call
# --------------------------------------------------------------------------

async def chunk_notes_async(
    course_id: str, lecture_id: str, source_text: str, source_type: str, lecture_date: str = None,
) -> dict:
    if not source_text.strip():
        raise MalformedSourceError("source has no extractable text to chunk")

    client = get_client()

    syllabus = await sync_to_async(storage.read_syllabus)(course_id)
    if syllabus is None:
        raise CourseNotFoundError(f"no syllabus.json found for course '{course_id}'")

    course_topics = syllabus.get("topics", [])

    user_prompt = (
        f"Course topics (reuse these exact strings when a chunk matches one):\n"
        f"{json.dumps(course_topics, indent=2)}\n\n"
        f"Lecture source text:\n---\n{source_text}\n---"
    )

    response = await client.messages.create(
        model=MODEL,
        max_tokens=4000,
        system=CHUNKING_SYSTEM_PROMPT,
        messages=[{"role": "user", "content": user_prompt}],
    )

    raw = "".join(block.text for block in response.content if block.type == "text").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw.strip())

    try:
        data = json.loads(raw)
    except json.JSONDecodeError as e:
        raise ValueError(f"model did not return valid JSON: {e}\n\nRaw output:\n{raw}")

    chunks = _dedupe_chunk_ids(data.get("chunks", []))

    return {
        "lecture_id": lecture_id,
        "source": source_type,
        "date": lecture_date,
        "topics": sorted({c["topic"] for c in chunks if c["topic"]}),
        "chunks": chunks,
    }
