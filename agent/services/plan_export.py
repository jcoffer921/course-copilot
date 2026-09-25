"""Render the current session-only academic plan as a DOCX download."""

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from django.utils.text import slugify


DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def plan_filename(draft_plan: dict) -> str:
    major = slugify(str(draft_plan.get("student_major") or "draft"))[:80] or "draft"
    return f"academic-plan-{major}.docx"


def render_plan_docx(draft_plan: dict) -> BytesIO:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)

    title = document.add_heading("Academic Plan", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    major = str(draft_plan.get("student_major") or "Undeclared")
    subtitle = document.add_paragraph(major)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(13)

    for semester in draft_plan.get("semesters") or []:
        document.add_heading(str(semester.get("label") or "Semester"), level=1)
        courses = semester.get("courses") or []
        if not courses:
            document.add_paragraph("No courses listed.")
            continue
        table = document.add_table(rows=1, cols=3)
        table.style = "Light Shading Accent 1"
        for cell, label in zip(table.rows[0].cells, ("Course", "Title", "Credits")):
            cell.text = label
        for course in courses:
            cells = table.add_row().cells
            cells[0].text = str(course.get("code") or "")
            cells[1].text = str(course.get("title") or "")
            cells[2].text = str(course.get("credits") or 0)

    document.add_paragraph()
    total = document.add_paragraph()
    total.add_run("Total planned credits: ").bold = True
    total.add_run(str(draft_plan.get("total_credits") or 0))

    notes = str(draft_plan.get("notes") or "").strip()
    if notes:
        document.add_heading("Notes", level=1)
        document.add_paragraph(notes)

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
