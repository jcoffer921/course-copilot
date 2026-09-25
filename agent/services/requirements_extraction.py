"""
Program-requirements Excel extraction logic — faculty upload a
requirements spreadsheet, this turns it into structured JSON for review.
Same async/JSON-envelope shape as syllabus_extraction.py, and reuses its
"extract only what's present, never infer" contract: this data seeds every
student-specific plan Cora drafts later, so a fabricated course or credit
count here would propagate into every plan built from it.

Nothing in this module writes to disk or the database — see Task 3's
import/confirm split, agent/services/requirements_extraction.py:
extract_requirements_async only ever returns JSON for the caller to review.
"""

import json
import re
from pathlib import Path

from .client import MODEL_DEFAULT as MODEL, create_message, get_client

REQUIREMENTS_EXTRACTION_SYSTEM_PROMPT = """You extract structured program \
requirements from a degree/major requirements spreadsheet.

Rules:
- Use ONLY information present in the provided spreadsheet text. Do not add a \
category, course, credit count, or prerequisite that is not explicitly stated \
in the source. Do not infer a course that isn't listed just because it's a \
common prerequisite for a course that IS listed.
- If a course row is missing its code, drop that course entry entirely rather \
than guessing a code or leaving it blank — a course without a code can't be \
matched against a transcript later. A course row missing only credits, \
prerequisites, or notes should still be emitted, with credits defaulting to 0 \
and prerequisites/notes empty, since those are commonly omitted in real \
spreadsheets.
- If a category's required credit total isn't stated anywhere, default \
credits_required to 0 rather than guessing.
- Output ONLY valid JSON matching the schema below. No preamble, no markdown \
fences, no commentary.

Schema:
{
  "program_name": "string",
  "catalog_year": "string",
  "total_credits_required": 0,
  "categories": [
    {
      "name": "string",
      "credits_required": 0,
      "courses": [
        {
          "code": "string",
          "title": "string",
          "credits": 0,
          "prerequisites": ["string"],
          "notes": "string"
        }
      ]
    }
  ]
}

Notes on fields:
- "program_name" and "catalog_year" should be read from a title/header row or \
sheet name if present. If genuinely absent from the document, use an empty \
string rather than guessing.
- "prerequisites" is a list of course codes as written in the source (e.g. \
"CS 101"). Leave it empty when no prerequisite is stated for that course.
- "notes" carries any qualifying text attached to a course row (e.g. "take in \
junior year", "may substitute with advisor approval"). Leave it an empty \
string when there's nothing to carry.
"""


class MalformedRequirementsSourceError(ValueError):
    """Raised when the uploaded workbook has nothing usable to extract —
    e.g. no sheets, or every sheet is empty."""


# --------------------------------------------------------------------------
# Excel -> text extraction (sync — file I/O, not worth async)
# --------------------------------------------------------------------------

def extract_text_from_workbook_bytes(data: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in (".xlsx", ".xlsm"):
        raise MalformedRequirementsSourceError(
            f"unsupported file type '{suffix}'. Use .xlsx or .xlsm."
        )

    from openpyxl import load_workbook
    import io

    try:
        workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as e:
        raise MalformedRequirementsSourceError(f"couldn't read Excel file: {e}") from e

    sheet_blocks = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            sheet_blocks.append(f"Sheet: {sheet.title}\n" + "\n".join(rows))

    text = "\n\n".join(sheet_blocks).strip()
    if not text:
        raise MalformedRequirementsSourceError("workbook has no readable rows")
    return text


def read_source_text_from_upload(upload) -> str:
    """For the DRF view — reads a Django UploadedFile."""
    return extract_text_from_workbook_bytes(upload.read(), upload.name)


def read_source_text_from_path(path: Path) -> str:
    """For the CLI harness — reads a local file."""
    return extract_text_from_workbook_bytes(path.read_bytes(), path.name)


# --------------------------------------------------------------------------
# Extraction call
# --------------------------------------------------------------------------

async def extract_requirements_async(source_bytes: bytes, filename: str, user=None) -> dict:
    source_text = extract_text_from_workbook_bytes(source_bytes, filename)

    client = get_client()
    user_prompt = f"Program requirements spreadsheet text:\n---\n{source_text}\n---"

    response = await create_message(
        client, user,
        model=MODEL,
        max_tokens=4000,
        system=REQUIREMENTS_EXTRACTION_SYSTEM_PROMPT,
        messages=[{"role": "user", "content": user_prompt}],
    )

    raw = "".join(block.text for block in response.content if block.type == "text").strip()
    return _normalize_extracted_requirements(_parse_model_json(raw))


def _normalize_extracted_requirements(data: dict) -> dict:
    """Coerce missing fields to sane defaults; never fabricate a category or
    course the model didn't already include."""
    if not isinstance(data, dict):
        raise MalformedRequirementsSourceError("model returned a non-object result")

    normalized = {
        "program_name": str(data.get("program_name") or ""),
        "catalog_year": str(data.get("catalog_year") or ""),
        "total_credits_required": _as_int(data.get("total_credits_required")),
        "categories": [],
    }

    categories = data.get("categories")
    if not isinstance(categories, list):
        return normalized

    for category in categories:
        if not isinstance(category, dict):
            continue
        courses = []
        for course in category.get("courses") or []:
            if not isinstance(course, dict):
                continue
            code = str(course.get("code") or "").strip()
            if not code:
                continue
            prerequisites = course.get("prerequisites")
            courses.append({
                "code": code,
                "title": str(course.get("title") or ""),
                "credits": _as_int(course.get("credits")),
                "prerequisites": [str(p) for p in prerequisites] if isinstance(prerequisites, list) else [],
                "notes": str(course.get("notes") or ""),
            })
        normalized["categories"].append({
            "name": str(category.get("name") or ""),
            "credits_required": _as_int(category.get("credits_required")),
            "courses": courses,
        })

    return normalized


def _as_int(value) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return 0


REQUIRED_TOP_LEVEL_KEYS = {"program_name", "catalog_year", "total_credits_required", "categories"}


def _parse_model_json(raw: str) -> dict:
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw.strip())

    try:
        return json.loads(raw)
    except json.JSONDecodeError as e:
        decoder = json.JSONDecoder()
        candidates = []
        for match in re.finditer(r"{", raw):
            try:
                obj, _ = decoder.raw_decode(raw[match.start():])
            except json.JSONDecodeError:
                continue
            if isinstance(obj, dict) and REQUIRED_TOP_LEVEL_KEYS.issubset(obj):
                candidates.append(obj)
        if candidates:
            return candidates[-1]
        raise ValueError(f"model did not return valid JSON: {e}\n\nRaw output:\n{raw}")
