"""
Syllabus extraction logic. Uses AsyncAnthropic so the DRF async view can
await it directly without blocking the ASGI event loop; the management
command wraps the same coroutine with asyncio.run() for standalone CLI use.
"""

import io
import json
import re
from datetime import datetime
from pathlib import Path

from . import storage
from .client import MODEL_DEFAULT as MODEL, get_client

EXTRACTION_SYSTEM_PROMPT = """You extract structured data from a course syllabus.

Rules:
- Use ONLY information present in the provided syllabus text. Do not add dates, \
topics, or grading components that are not explicitly stated or clearly implied \
by the document (e.g. "Midterm — Oct 14" clearly implies an exam on that date).
- If any required field of a dates[] or grading[] entry can't be determined, drop \
the ENTIRE entry rather than guessing or emitting it with that field missing — \
every dates[] item needs date AND title AND type together, every grading[] item \
needs component AND weight_pct together. A partial object (e.g. title/type with \
no date) is invalid output; omitting the whole entry is the correct move instead. \
Do not invent plausible-sounding data.
- Output ONLY valid JSON matching the schema below. No preamble, no markdown \
fences, no commentary.

Schema:
{
  "course_id": "string",
  "course_name": "string",
  "meeting_patterns": [{"title": "string", "days": [0], "start_time": "HH:MM", "end_time": "HH:MM|null", "start_date": "YYYY-MM-DD|null", "end_date": "YYYY-MM-DD|null"}],
  "dates": [{"date": "YYYY-MM-DD", "title": "string", "type": "hw|project|test_quiz|class|other"}],
  "grading": [{"component": "Homework|Tests|Quizzes|Midterm|Final|Projects|Lab and Demo|Final Project|Class Participation|Other", "weight_pct": 0}],
  "topics": ["string"]
}

Notes on fields:
- "meeting_patterns" contains recurring class meeting information only when it is explicitly present. Use weekday numbers Monday=0 through Sunday=6. Preserve explicit start/end dates when stated. Do not infer a current-semester schedule from an old syllabus or from the course's total weekly duration. Return an empty list when days or start time are not stated.
- "date" must be YYYY-MM-DD, and the year is NOT optional. Before extracting any \
dates[] entries, check whether an explicit calendar year (e.g. "2026") appears \
ANYWHERE in the document — in a header, term label ("Fall 2026"), or elsewhere. \
If no explicit year appears anywhere in the document, you cannot determine the \
year for ANY date in it: the dates[] array must be empty, even though the \
document clearly has months/days/weeks in it. A vague term reference like "this \
semester", "this term", or "the current semester" is NOT a year and is NOT a \
license to guess one — guessing a plausible-looking year (e.g. defaulting to the \
current year) is fabrication and is exactly what this rule exists to prevent. \
Only extract a date when the year is either stated directly on it or unambiguous \
from an explicit year stated elsewhere in the same document.
- "type" must be exactly one of: hw, project, test_quiz, class, other. Map the \
syllabus's own wording onto these:
    - hw: homework, assignments, problem sets, exercises
    - project: projects, presentations, capstone work, project deliverables
    - test_quiz: quizzes, tests, exams, midterms, finals
    - class: lectures, readings, class meetings, labs, demos
    - other: anything real that doesn't fit the categories above
- "weight_pct" is a number (e.g. 20 for 20%). If weights aren't given, drop that \
whole grading[] entry — never emit {"component": ...} with no "weight_pct".
- "component" MUST be exactly one of: Homework, Tests, Quizzes, Midterm, Final, \
Projects, Lab and Demo, Final Project, Class Participation, Other. Map the \
syllabus's own wording onto these:
    - Homework: homework, assignments, problem sets, exercises
    - Tests: recurring or unlabeled tests/exams not specifically called out as \
"the midterm" or "the final"
    - Quizzes: quizzes
    - Midterm: an exam explicitly labeled as the midterm
    - Final: an exam explicitly labeled as the final
    - Projects: projects, presentations, capstone work that is not explicitly \
the final project
    - Lab and Demo: lab work, labs, demos, lab demonstrations
    - Final Project: an explicitly labeled final project
    - Class Participation: class participation, participation
    - Other: anything real that doesn't fit the categories above (attendance, \
miscellaneous course requirements, etc.) — use Other rather than dropping the \
entry or forcing it into the wrong bucket.
  If two or more syllabus lines map to the same bucket (e.g. "Problem Sets" 15% \
and "Weekly Exercises" 10%, both Homework), merge them into ONE grading[] entry \
with the summed weight_pct — never emit two entries with the same component.
- "topics" is a flat list of topic/unit names as they appear in the syllabus.
"""


# --------------------------------------------------------------------------
# Text extraction from source bytes (sync — file I/O, not worth async)
# --------------------------------------------------------------------------

SUPPORTED_SOURCE_TYPES = ".pdf, .pptx, .docx, .txt, or .md"


def _extract_docx_text(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise ValueError("couldn't read .docx file: python-docx is not installed") from e

    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        raise ValueError(f"couldn't read .docx file: {e}") from e

    parts = []
    parts.extend(p.text.strip() for p in doc.paragraphs if p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n".join(parts).strip()
    if not text:
        raise ValueError("no extractable text found in Word document")
    return text


def _extract_pptx_text(data: bytes) -> str:
    try:
        from pptx import Presentation
    except ImportError as e:
        raise ValueError("couldn't read .pptx file: python-pptx is not installed") from e

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as e:
        raise ValueError(f"couldn't read .pptx file: {e}") from e

    parts = []
    for slide_number, slide in enumerate(prs.slides, start=1):
        slide_parts = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if text:
                slide_parts.append(text)
        if slide_parts:
            parts.append(f"Slide {slide_number}\n" + "\n".join(slide_parts))

    text = "\n\n".join(parts).strip()
    if not text:
        raise ValueError("no extractable text found in PowerPoint deck")
    return text


def extract_text_from_bytes(data: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()

    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages).strip()
        if not text:
            raise ValueError(
                "no extractable text found in PDF (it may be a scanned image — "
                "OCR is not supported yet)"
            )
        return text

    if suffix == ".pptx":
        return _extract_pptx_text(data)

    if suffix == ".docx":
        return _extract_docx_text(data)

    if suffix in (".txt", ".md"):
        text = data.decode("utf-8", errors="ignore").strip()
        if not text:
            raise ValueError("source file is empty")
        return text

    raise ValueError(f"unsupported file type '{suffix}'. Use {SUPPORTED_SOURCE_TYPES}.")


def read_source_text_from_path(path: Path) -> str:
    """For the CLI management command — reads a local file."""
    return extract_text_from_bytes(path.read_bytes(), path.name)


def read_source_text_from_upload(upload) -> str:
    """For the DRF view — reads a Django UploadedFile."""
    return extract_text_from_bytes(upload.read(), upload.name)


# --------------------------------------------------------------------------
# Extraction call
# --------------------------------------------------------------------------

async def extract_syllabus_async(source_text: str, course_id: str, course_name_hint: str = None) -> dict:
    client = get_client()

    user_prompt = f"course_id to use: {course_id}\n"
    if course_name_hint:
        user_prompt += f"course_name hint (use if it matches the document): {course_name_hint}\n"
    user_prompt += f"\nSyllabus text:\n---\n{source_text}\n---"

    response = await client.messages.create(
        model=MODEL,
        max_tokens=4000,
        system=EXTRACTION_SYSTEM_PROMPT,
        messages=[{"role": "user", "content": user_prompt}],
    )

    raw = "".join(block.text for block in response.content if block.type == "text").strip()
    normalized = _normalize_extracted_syllabus(_parse_model_json(raw))
    source_patterns = _extract_meeting_patterns(source_text)
    if source_patterns:
        normalized["meeting_patterns"] = source_patterns
    return normalized


def _normalize_clock_time(value: str) -> str | None:
    text = re.sub(r"[.\s]", "", value or "").lower()
    match = re.fullmatch(r"(\d{1,2})(?::(\d{2}))?(am|pm)?", text)
    if not match:
        return None
    hour, minute, period = int(match.group(1)), int(match.group(2) or 0), match.group(3)
    if minute > 59 or (period and not 1 <= hour <= 12) or (not period and hour > 23):
        return None
    if period:
        hour = hour % 12 + (12 if period == "pm" else 0)
    return f"{hour:02d}:{minute:02d}"


def _extract_meeting_patterns(source_text: str) -> list[dict]:
    """Extract explicit recurring class hours that the broad model schema may omit."""
    weekday_words = {
        "mon": 0, "tue": 1, "wed": 2, "thu": 3, "fri": 4, "sat": 5, "sun": 6,
    }
    day_re = re.compile(
        r"\b(mon(?:day)?s?|tue(?:sday)?s?|wed(?:nesday)?s?|thu(?:rsday)?s?|fri(?:day)?s?|sat(?:urday)?s?|sun(?:day)?s?)\b",
        re.I,
    )
    time_re = re.compile(
        r"\b(?:(?:1[0-2]|0?[1-9])(?::[0-5]\d)?\s*(?:a\.?m\.?|p\.?m\.?)|(?:[01]?\d|2[0-3]):[0-5]\d)\b",
        re.I,
    )
    date_range = re.search(
        r"\b([A-Z][a-z]{2,8}\s+\d{1,2},\s*\d{4})\s*[-–—]\s*([A-Z][a-z]{2,8}\s+\d{1,2},\s*\d{4})\b",
        source_text or "",
    )
    start_date = end_date = None
    if date_range:
        for fmt in ("%b %d, %Y", "%B %d, %Y"):
            try:
                start_date = datetime.strptime(date_range.group(1), fmt).date().isoformat()
                end_date = datetime.strptime(date_range.group(2), fmt).date().isoformat()
                break
            except ValueError:
                continue

    patterns = []
    for line in (source_text or "").splitlines():
        if not re.search(r"\b(class|course)\s+(?:meeting\s+)?(?:hours?|times?)\b|\bmeets?\b", line, re.I):
            continue
        days = {weekday_words[m.group(1).lower()[:3]] for m in day_re.finditer(line)}
        compact = re.search(r"\b(MoWeFr|MWF|TTh|TuTh)\b", line, re.I)
        if compact:
            days.update({0, 2, 4} if compact.group(0).lower() in {"mowefr", "mwf"} else {1, 3})
        times = [_normalize_clock_time(match.group(0)) for match in time_re.finditer(line)]
        times = [value for value in times if value]
        if not days or not times:
            continue
        patterns.append({
            "title": "Class",
            "days": sorted(days),
            "start_time": times[0],
            "end_time": times[1] if len(times) > 1 else None,
            "start_date": start_date,
            "end_date": end_date,
        })
    return patterns


def _normalize_extracted_syllabus(data: dict) -> dict:
    """Bring model output forward to the app's current schema.

    Older prompts and occasional model responses may use legacy calendar
    types such as "exam" or "assignment"; the rest of the app now stores
    hw/project/test_quiz/class/other.
    """
    if not isinstance(data, dict):
        return data
    data = dict(data)
    data.setdefault("meeting_patterns", [])
    dates = data.get("dates")
    if isinstance(dates, list):
        normalized_dates = []
        for item in dates:
            if isinstance(item, dict):
                normalized_dates.append(dict(item, type=storage.normalize_date_type(item.get("type"))))
            else:
                normalized_dates.append(item)
        data["dates"] = normalized_dates
    return data


def _parse_model_json(raw: str) -> dict:
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw.strip())

    try:
        return json.loads(raw)
    except json.JSONDecodeError as e:
        # The model sometimes prefixes/trails prose, or emits an initial JSON
        # object followed by a self-corrected second object. Decode complete JSON
        # objects we can find and use the last one that has the syllabus top-level
        # shape, which preserves the common self-correction case without
        # accidentally returning a nested grading/date object.
        decoder = json.JSONDecoder()
        candidates = []
        required_keys = {"course_id", "course_name", "dates", "grading", "topics"}
        for match in re.finditer(r"{", raw):
            try:
                obj, _ = decoder.raw_decode(raw[match.start():])
            except json.JSONDecodeError:
                continue
            if isinstance(obj, dict) and required_keys.issubset(obj):
                candidates.append(obj)
        if candidates:
            return candidates[-1]
        raise ValueError(f"model did not return valid JSON: {e}\n\nRaw output:\n{raw}")
