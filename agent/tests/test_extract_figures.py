import io

import pymupdf
import pytest
from pptx import Presentation
from pptx.util import Inches

from agent.services import extract_figures, storage


@pytest.fixture
def isolated_courses_dir(tmp_path, monkeypatch):
    """Redirects storage.COURSES_DIR to a throwaway tmp_path so these tests
    never touch real course data."""
    monkeypatch.setattr(storage, "COURSES_DIR", tmp_path)
    return tmp_path


def _solid_pixmap(size: int, rgb: tuple) -> pymupdf.Pixmap:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, size, size))
    pix.set_rect(pix.irect, rgb)
    return pix


def _write_notes(course_id, lecture_id, chunks, user):
    storage.write_notes(course_id, lecture_id, {
        "lecture_id": lecture_id,
        "source": "notes",
        "date": "2026-08-25",
        "topics": sorted({c["topic"] for c in chunks}),
        "chunks": chunks,
    }, user)


# --------------------------------------------------------------------------
# PDF extraction: embedded images, rasterization fallback, junk filtering
# --------------------------------------------------------------------------

def _build_three_page_pdf(tmp_path):
    """Page 1: real text + a large (200x200) embedded image.
    Page 2: real text + a vector-drawn rectangle, no embedded image object.
    Page 3: real text + a small (40x40) embedded logo-sized image."""
    doc = pymupdf.open()

    page1 = doc.new_page()
    page1.insert_text((50, 72), "Recursion base case stops the recursive calls immediately.")
    page1.insert_image(pymupdf.Rect(300, 300, 500, 500), pixmap=_solid_pixmap(200, (200, 50, 50)))

    page2 = doc.new_page()
    page2.insert_text((50, 72), "Sorting algorithms like quicksort partition the array recursively.")
    page2.draw_rect(pymupdf.Rect(50, 100, 300, 300), color=(0, 0, 1), fill=(0, 0, 1))

    page3 = doc.new_page()
    page3.insert_text((50, 72), "Course logistics and syllabus reminders for this week.")
    page3.insert_image(pymupdf.Rect(10, 10, 50, 50), pixmap=_solid_pixmap(40, (10, 10, 10)))

    path = tmp_path / "lecture.pdf"
    doc.save(path)
    doc.close()
    return path


def test_extract_pdf_returns_one_candidate_per_page_by_kind(tmp_path):
    path = _build_three_page_pdf(tmp_path)

    candidates, page_texts = extract_figures._extract_pdf(path)

    by_page = {c["page"]: c for c in candidates}
    assert by_page[1]["kind"] == "embedded"
    assert by_page[1]["width"] == 200 and by_page[1]["height"] == 200
    assert by_page[2]["kind"] == "rasterized_page"  # forced by page 2 having no embedded image object
    assert by_page[2]["width"] > extract_figures.MIN_IMAGE_DIMENSION_PX
    assert by_page[3]["kind"] == "embedded"
    assert by_page[3]["width"] == 40 and by_page[3]["height"] == 40
    assert "Recursion base case" in page_texts[1]
    assert "quicksort" in page_texts[2]


def test_filter_junk_drops_small_embedded_but_keeps_rasterized_page(tmp_path):
    path = _build_three_page_pdf(tmp_path)
    candidates, _ = extract_figures._extract_pdf(path)

    kept = extract_figures._filter_junk(candidates)

    kept_by_page = {c["page"]: c["kind"] for c in kept}
    assert kept_by_page == {1: "embedded", 2: "rasterized_page"}
    assert 3 not in kept_by_page  # the 40x40 logo-sized image is junk


def test_extract_pdf_raises_source_file_error_on_malformed_pdf(tmp_path):
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"this is not a real pdf")

    with pytest.raises(extract_figures.SourceFileError):
        extract_figures._extract_pdf(bad)


def test_extract_pdf_raises_source_file_error_on_encrypted_pdf(tmp_path):
    doc = pymupdf.open()
    doc.new_page().insert_text((50, 72), "secret")
    path = tmp_path / "locked.pdf"
    doc.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="secret", owner_pw="secret")
    doc.close()

    with pytest.raises(extract_figures.SourceFileError):
        extract_figures._extract_pdf(path)


# --------------------------------------------------------------------------
# PPTX extraction: embedded picture shapes, multiple images per slide
# --------------------------------------------------------------------------

def _build_two_slide_pptx(tmp_path):
    """Slide 1: two embedded images (one real-sized, one logo-sized) plus text.
    Slide 2: one real-sized embedded image plus different-topic text."""
    prs = Presentation()

    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    slide1.shapes.add_textbox(Inches(1), Inches(0.2), Inches(8), Inches(1)).text_frame.text = (
        "Recursion base case stops the recursive calls."
    )
    slide1.shapes.add_picture(
        io.BytesIO(_solid_pixmap(200, (200, 50, 50)).tobytes("png")), Inches(1), Inches(1), width=Inches(2), height=Inches(2),
    )
    slide1.shapes.add_picture(
        io.BytesIO(_solid_pixmap(30, (10, 10, 10)).tobytes("png")), Inches(4), Inches(1), width=Inches(0.3), height=Inches(0.3),
    )

    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    slide2.shapes.add_textbox(Inches(1), Inches(0.2), Inches(8), Inches(1)).text_frame.text = (
        "Sorting algorithms partition arrays recursively."
    )
    slide2.shapes.add_picture(
        io.BytesIO(_solid_pixmap(180, (50, 200, 50)).tobytes("png")), Inches(1), Inches(1), width=Inches(2), height=Inches(2),
    )

    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def test_extract_pptx_returns_every_embedded_picture_shape(tmp_path):
    path = _build_two_slide_pptx(tmp_path)

    candidates, page_texts = extract_figures._extract_pptx(path)

    slide1_images = [c for c in candidates if c["page"] == 1]
    slide2_images = [c for c in candidates if c["page"] == 2]
    assert len(slide1_images) == 2
    assert len(slide2_images) == 1
    assert "Recursion base case" in page_texts[1]
    assert "Sorting algorithms" in page_texts[2]


def test_filter_junk_on_pptx_drops_logo_sized_embedded_image(tmp_path):
    path = _build_two_slide_pptx(tmp_path)
    candidates, _ = extract_figures._extract_pptx(path)

    kept = extract_figures._filter_junk(candidates)

    assert len(kept) == 2  # the 30x30 logo on slide 1 is dropped, the two real images survive
    assert all(c["width"] >= extract_figures.MIN_IMAGE_DIMENSION_PX for c in kept)


def test_extract_pptx_raises_source_file_error_on_missing_file(tmp_path):
    with pytest.raises(extract_figures.SourceFileError):
        extract_figures._extract_pptx(tmp_path / "does-not-exist.pptx")


def test_extract_pptx_raises_source_file_error_on_malformed_file(tmp_path):
    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"not a real pptx")

    with pytest.raises(extract_figures.SourceFileError):
        extract_figures._extract_pptx(bad)


# --------------------------------------------------------------------------
# Proximity chunk-linking: deterministic, documented, and never a blind guess
# --------------------------------------------------------------------------

def test_resolve_chunk_pages_links_each_chunk_to_its_best_overlapping_page():
    chunks = [
        {"id": "recursion-base", "text": "Recursion base case stops the recursive calls immediately."},
        {"id": "quicksort", "text": "Sorting algorithms like quicksort partition the array recursively."},
    ]
    page_texts = {
        1: "Recursion base case stops the recursive calls immediately.",
        2: "Sorting algorithms like quicksort partition the array recursively.",
        3: "Course logistics and syllabus reminders for this week.",
    }

    resolved = extract_figures._resolve_chunk_pages(chunks, page_texts)

    assert resolved == {"recursion-base": 1, "quicksort": 2}


def test_resolve_chunk_pages_breaks_ties_deterministically_toward_lowest_page():
    chunks = [{"id": "shared", "text": "recursion recursion recursion"}]
    page_texts = {3: "recursion appears here too", 1: "recursion appears here too"}

    resolved = extract_figures._resolve_chunk_pages(chunks, page_texts)

    assert resolved == {"shared": 1}  # equal overlap on page 1 and 3 — lowest page wins, not insertion order


def test_resolve_chunk_pages_leaves_chunk_unresolved_when_no_overlap():
    chunks = [{"id": "unrelated", "text": "xyzxyz qqqqq zzzzzz"}]
    page_texts = {1: "recursion base case", 2: "sorting quicksort partition"}

    resolved = extract_figures._resolve_chunk_pages(chunks, page_texts)

    assert resolved == {}  # never guess a page with zero real signal


# --------------------------------------------------------------------------
# build_extraction_plan / commit_extraction_plan — end to end, plan-then-pause,
# and the required failure modes (missing source, malformed source, no notes)
# --------------------------------------------------------------------------

@pytest.mark.django_db
def test_build_extraction_plan_raises_on_missing_source_file(isolated_courses_dir, django_user_model):
    user = django_user_model.objects.create_user(username="ef-missing-src", email="ef-missing-src@example.com")
    storage.write_syllabus("cs101", {"course_id": "cs101", "course_name": "CS101", "dates": [], "grading": [], "topics": []}, user)
    _write_notes("cs101", "lecture01", [{"id": "c1", "topic": "t", "text": "text"}], user)

    with pytest.raises(extract_figures.SourceFileError):
        extract_figures.build_extraction_plan("cs101", "lecture01", isolated_courses_dir / "nope.pdf", user)


@pytest.mark.django_db
def test_build_extraction_plan_raises_on_unsupported_file_type(isolated_courses_dir, django_user_model, tmp_path):
    user = django_user_model.objects.create_user(username="ef-bad-type", email="ef-bad-type@example.com")
    storage.write_syllabus("cs101", {"course_id": "cs101", "course_name": "CS101", "dates": [], "grading": [], "topics": []}, user)
    _write_notes("cs101", "lecture01", [{"id": "c1", "topic": "t", "text": "text"}], user)
    txt_path = tmp_path / "notes.txt"
    txt_path.write_text("plain text notes", encoding="utf-8")

    with pytest.raises(extract_figures.SourceFileError):
        extract_figures.build_extraction_plan("cs101", "lecture01", txt_path, user)


@pytest.mark.django_db
def test_build_extraction_plan_raises_on_malformed_pdf(isolated_courses_dir, django_user_model, tmp_path):
    user = django_user_model.objects.create_user(username="ef-malformed", email="ef-malformed@example.com")
    storage.write_syllabus("cs101", {"course_id": "cs101", "course_name": "CS101", "dates": [], "grading": [], "topics": []}, user)
    _write_notes("cs101", "lecture01", [{"id": "c1", "topic": "t", "text": "text"}], user)
    bad_pdf = tmp_path / "bad.pdf"
    bad_pdf.write_bytes(b"not a real pdf")

    with pytest.raises(extract_figures.SourceFileError):
        extract_figures.build_extraction_plan("cs101", "lecture01", bad_pdf, user)


@pytest.mark.django_db
def test_build_extraction_plan_raises_when_lecture_notes_do_not_exist_yet(isolated_courses_dir, django_user_model, tmp_path):
    user = django_user_model.objects.create_user(username="ef-no-notes", email="ef-no-notes@example.com")
    storage.write_syllabus("cs101", {"course_id": "cs101", "course_name": "CS101", "dates": [], "grading": [], "topics": []}, user)
    pdf_path = _build_three_page_pdf(tmp_path)

    with pytest.raises(extract_figures.LectureNotFoundError):
        extract_figures.build_extraction_plan("cs101", "lecture01", pdf_path, user)


@pytest.mark.django_db
def test_build_and_commit_extraction_plan_links_images_to_chunks_by_page(isolated_courses_dir, django_user_model, tmp_path):
    """Full happy path over the 3-page fixture: page 1's image links to the
    recursion chunk, page 2's rasterized fallback links to the sorting
    chunk, page 3's logo is filtered out entirely and never written."""
    user = django_user_model.objects.create_user(username="ef-e2e", email="ef-e2e@example.com")
    storage.write_syllabus("cs101", {
        "course_id": "cs101", "course_name": "CS101", "dates": [], "grading": [], "topics": ["Recursion", "Sorting"],
    }, user)
    _write_notes("cs101", "lecture01", [
        {"id": "recursion-base", "topic": "Recursion", "text": "Recursion base case stops the recursive calls immediately."},
        {"id": "quicksort", "topic": "Sorting", "text": "Sorting algorithms like quicksort partition the array recursively."},
    ], user)
    pdf_path = _build_three_page_pdf(tmp_path)

    plan = extract_figures.build_extraction_plan("cs101", "lecture01", pdf_path, user)
    assert len(plan["images"]) == 2  # the page-3 logo never becomes a candidate worth writing
    assert plan["existing_manifest"] is None

    result = extract_figures.commit_extraction_plan(plan, user, overwrite=False)
    assert result == {
        "image_count": 2, "embedded_count": 1, "rasterized_count": 1,
        "linked_chunk_count": 2, "skipped_undecodable": 0,
    }

    manifest = storage.read_image_manifest("cs101", "lecture01", user)
    manifest_by_page = {entry["source_page"]: entry for entry in manifest}
    assert manifest_by_page[1]["chunk_ids"] == ["recursion-base"]
    assert manifest_by_page[1]["match_method"] == "proximity"
    assert manifest_by_page[2]["chunk_ids"] == ["quicksort"]
    assert {entry["source_page"] for entry in manifest} == {1, 2}  # page 3 wrote nothing

    note = storage.read_lecture("cs101", "lecture01", user)
    chunks_by_id = {c["id"]: c for c in note["chunks"]}
    assert chunks_by_id["recursion-base"]["page"] == 1
    assert chunks_by_id["recursion-base"]["image_ids"] == [manifest_by_page[1]["image_id"]]
    assert chunks_by_id["quicksort"]["page"] == 2
    assert chunks_by_id["quicksort"]["image_ids"] == [manifest_by_page[2]["image_id"]]
    # untouched fields survive the additive update
    assert chunks_by_id["recursion-base"]["topic"] == "Recursion"
    assert chunks_by_id["recursion-base"]["text"] == "Recursion base case stops the recursive calls immediately."

    for entry in manifest:
        path = storage.lecture_image_path("cs101", "lecture01", entry["image_id"], user)
        assert path.exists()
        assert path.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n"  # real PNG bytes, never a placeholder


@pytest.mark.django_db
def test_multiple_images_and_multiple_chunks_on_the_same_page_link_deterministically(
    isolated_courses_dir, django_user_model, tmp_path,
):
    """Two chunks whose text both best-matches page 1, and page 1 has two
    embedded images — every image on that page must link to both chunks, in
    the lecture's original chunk order, not an arbitrary/unstable order."""
    user = django_user_model.objects.create_user(username="ef-multi", email="ef-multi@example.com")
    storage.write_syllabus("cs101", {
        "course_id": "cs101", "course_name": "CS101", "dates": [], "grading": [], "topics": ["Recursion"],
    }, user)
    _write_notes("cs101", "lecture01", [
        {"id": "chunk-a", "topic": "Recursion", "text": "Recursion base case example diagram walkthrough."},
        {"id": "chunk-b", "topic": "Recursion", "text": "Recursion base case second explanation of the same diagram."},
    ], user)

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Recursion base case example diagram walkthrough explanation.")
    page.insert_image(pymupdf.Rect(50, 300, 250, 500), pixmap=_solid_pixmap(200, (200, 50, 50)))
    page.insert_image(pymupdf.Rect(300, 300, 500, 500), pixmap=_solid_pixmap(200, (50, 200, 50)))
    pdf_path = tmp_path / "lecture.pdf"
    doc.save(pdf_path)
    doc.close()

    plan = extract_figures.build_extraction_plan("cs101", "lecture01", pdf_path, user)
    result = extract_figures.commit_extraction_plan(plan, user, overwrite=False)

    assert result["image_count"] == 2
    manifest = storage.read_image_manifest("cs101", "lecture01", user)
    assert len(manifest) == 2
    for entry in manifest:
        assert entry["chunk_ids"] == ["chunk-a", "chunk-b"]  # deterministic: original chunk order, both linked

    note = storage.read_lecture("cs101", "lecture01", user)
    chunks_by_id = {c["id"]: c for c in note["chunks"]}
    assert set(chunks_by_id["chunk-a"]["image_ids"]) == {e["image_id"] for e in manifest}
    assert set(chunks_by_id["chunk-b"]["image_ids"]) == {e["image_id"] for e in manifest}


@pytest.mark.django_db
def test_commit_extraction_plan_raises_file_exists_without_overwrite(isolated_courses_dir, django_user_model, tmp_path):
    user = django_user_model.objects.create_user(username="ef-conflict", email="ef-conflict@example.com")
    storage.write_syllabus("cs101", {"course_id": "cs101", "course_name": "CS101", "dates": [], "grading": [], "topics": []}, user)
    _write_notes("cs101", "lecture01", [
        {"id": "recursion-base", "topic": "Recursion", "text": "Recursion base case stops the recursive calls immediately."},
    ], user)
    pdf_path = _build_three_page_pdf(tmp_path)

    plan = extract_figures.build_extraction_plan("cs101", "lecture01", pdf_path, user)
    extract_figures.commit_extraction_plan(plan, user, overwrite=False)

    plan2 = extract_figures.build_extraction_plan("cs101", "lecture01", pdf_path, user)
    assert plan2["existing_manifest"] is not None
    with pytest.raises(FileExistsError):
        extract_figures.commit_extraction_plan(plan2, user, overwrite=False)


@pytest.mark.django_db
def test_commit_extraction_plan_overwrite_replaces_existing_manifest_and_images(
    isolated_courses_dir, django_user_model, tmp_path,
):
    user = django_user_model.objects.create_user(username="ef-overwrite", email="ef-overwrite@example.com")
    storage.write_syllabus("cs101", {
        "course_id": "cs101", "course_name": "CS101", "dates": [], "grading": [], "topics": ["Recursion"],
    }, user)
    _write_notes("cs101", "lecture01", [
        {"id": "recursion-base", "topic": "Recursion", "text": "Recursion base case stops the recursive calls immediately."},
    ], user)
    pdf_path = _build_three_page_pdf(tmp_path)

    plan1 = extract_figures.build_extraction_plan("cs101", "lecture01", pdf_path, user)
    extract_figures.commit_extraction_plan(plan1, user, overwrite=False)
    first_manifest = storage.read_image_manifest("cs101", "lecture01", user)
    first_image_path = storage.lecture_image_path("cs101", "lecture01", first_manifest[0]["image_id"], user)

    plan2 = extract_figures.build_extraction_plan("cs101", "lecture01", pdf_path, user)
    result = extract_figures.commit_extraction_plan(plan2, user, overwrite=True)

    assert result["image_count"] == 2
    assert first_image_path.exists()  # same image_ids get regenerated identically from the same source


def test_normalize_to_png_returns_none_for_undecodable_bytes():
    assert extract_figures._normalize_to_png(b"not an image at all") is None


def test_normalize_to_png_converts_non_png_bytes_to_real_png():
    jpeg_bytes = _solid_pixmap(100, (10, 20, 30)).tobytes("jpg")

    png_bytes = extract_figures._normalize_to_png(jpeg_bytes)

    assert png_bytes[:8] == b"\x89PNG\r\n\x1a\n"
