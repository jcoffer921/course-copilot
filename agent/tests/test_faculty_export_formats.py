from io import BytesIO

import pytest
from docx import Document

from agent.services import plan_export, plan_state

LAYOUTS = ("term_table", "year_grid", "plan_notes")


def _draft(semesters=None, notes="Confirm elective availability."):
    if semesters is None:
        semesters = [
            {"label": "Fall 2026", "courses": [
                {"code": "CS 101", "title": "Intro to Programming", "credits": 3},
                {"code": "MATH 230", "title": "Discrete Mathematics", "credits": 3},
            ]},
            {"label": "Spring 2027", "courses": [{"code": "CS 201", "title": "Data Structures", "credits": 4}]},
        ]
    return {
        "student_major": "Computer Science",
        "semesters": semesters,
        "total_credits": 10,
        "notes": notes,
    }


def _open(buffer):
    return Document(BytesIO(buffer.getvalue()))


def _text(document):
    paragraphs = [p.text for p in document.paragraphs]
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
    return paragraphs, tables


def test_term_table_matches_existing_output():
    default = _text(_open(plan_export.render_plan_docx(_draft())))
    explicit = _text(_open(plan_export.render_plan_docx(_draft(), "term_table")))

    assert default == explicit
    paragraphs, tables = default
    assert len(tables) == 2
    assert tables[0][0] == ["Course", "Title", "Credits"]
    assert tables[0][1] == ["CS 101", "Intro to Programming", "3"]
    assert "Total planned credits: 10" in paragraphs
    assert "Notes" in paragraphs


def test_year_grid_places_semesters_in_one_table():
    _, tables = _text(_open(plan_export.render_plan_docx(_draft(), "year_grid")))

    assert len(tables) == 1
    assert tables[0][0] == ["Fall 2026", "Spring 2027"]
    assert "CS 101" in tables[0][1][0] and "MATH 230" in tables[0][1][0]
    assert "CS 201" in tables[0][1][1]


def test_year_grid_starts_a_new_table_after_three_semesters():
    semesters = [{"label": f"Term {i}", "courses": []} for i in range(1, 5)]

    _, tables = _text(_open(plan_export.render_plan_docx(_draft(semesters), "year_grid")))

    assert [len(table[0]) for table in tables] == [3, 1]


def test_plan_notes_includes_followup_section():
    paragraphs, _ = _text(_open(plan_export.render_plan_docx(_draft(), "plan_notes")))

    assert "Flagged issues and follow-ups" in paragraphs
    assert "Action items" in paragraphs
    assert "Confirm elective availability." in paragraphs


def test_unknown_layout_falls_back_to_term_table():
    assert _text(_open(plan_export.render_plan_docx(_draft(), "no-such-layout"))) == _text(
        _open(plan_export.render_plan_docx(_draft(), "term_table"))
    )


@pytest.mark.parametrize("layout", LAYOUTS)
def test_empty_semesters_still_produces_valid_docx(layout):
    document = _open(plan_export.render_plan_docx(_draft([], notes=""), layout))

    assert "Academic Plan" in [p.text for p in document.paragraphs]


def test_layout_constants():
    assert plan_export.EXPORT_LAYOUTS == LAYOUTS
    assert plan_export.DEFAULT_LAYOUT == "term_table"


# ---- Task 10: export endpoint `format` option ----

from rest_framework.test import APIClient

from agent.models import UserSettings

EXPORT_URL = "/api/faculty/plan/export/"


def _faculty_client(django_user_model, with_draft=True):
    user = django_user_model.objects.create_user(username="export-format-faculty")
    UserSettings.objects.create(user=user, access_status=UserSettings.ACCESS_ACTIVE, role=UserSettings.ROLE_FACULTY)
    client = APIClient()
    client.force_authenticate(user=user)
    if with_draft:
        plan_state.write(user, {
            "program_requirement_id": "00000000-0000-0000-0000-000000000000",
            "student_details": {"name": "Hidden Student"},
            "conversation": [{"role": "user", "content": "private conversation text"}],
            "draft_plan": _draft(),
        })
    return client


def _exported_document(response):
    return Document(BytesIO(b"".join(response.streaming_content)))


@pytest.mark.django_db
def test_export_defaults_to_term_table_without_body(django_user_model):
    response = _faculty_client(django_user_model).post(EXPORT_URL)

    assert response.status_code == 200
    assert len(_exported_document(response).tables) == 2


@pytest.mark.django_db
def test_export_honors_year_grid(django_user_model):
    response = _faculty_client(django_user_model).post(EXPORT_URL, {"format": "year_grid"}, format="json")

    assert response.status_code == 200
    assert len(_exported_document(response).tables) == 1


@pytest.mark.django_db
@pytest.mark.parametrize("value", ["bogus", "", None, 7])
def test_export_unknown_format_falls_back_and_returns_200(django_user_model, value):
    response = _faculty_client(django_user_model).post(EXPORT_URL, {"format": value}, format="json")

    assert response.status_code == 200
    assert len(_exported_document(response).tables) == 2


@pytest.mark.django_db
def test_export_without_draft_still_400(django_user_model):
    response = _faculty_client(django_user_model, with_draft=False).post(EXPORT_URL, {"format": "year_grid"}, format="json")

    assert response.status_code == 400
    assert response.data["detail"] == "Create a draft plan before downloading it."


@pytest.mark.django_db
@pytest.mark.parametrize("layout", LAYOUTS)
def test_export_contains_no_session_metadata(django_user_model, layout):
    response = _faculty_client(django_user_model).post(EXPORT_URL, {"format": layout}, format="json")

    document = _exported_document(response)
    everything = "\n".join(p.text for p in document.paragraphs) + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    assert "private conversation text" not in everything
    assert "Hidden Student" not in everything


# ---- Final review fixes: grounded export values and session-mismatch guard ----

import uuid

from agent.models import ProgramRequirement


def _client_with_program_session(django_user_model, username="export-grounding-faculty"):
    user = django_user_model.objects.create_user(username=username)
    UserSettings.objects.create(user=user, access_status=UserSettings.ACCESS_ACTIVE, role=UserSettings.ROLE_FACULTY)
    program = ProgramRequirement.objects.create(
        user=user, program_name="B.S. Test", catalog_year="2026",
        requirements={"program_name": "B.S. Test", "catalog_year": "2026", "total_credits_required": 120, "categories": [
            {"name": "Core", "credits_required": 7, "courses": [
                {"code": "CS 201", "title": "Data Structures", "credits": 3, "prerequisites": [], "notes": ""},
            ]},
        ]},
    )
    client = APIClient()
    client.force_authenticate(user=user)
    plan_state.write(user, {
        "program_requirement_id": str(program.requirement_id),
        "student_details": {}, "conversation": [],
        "draft_plan": {
            "student_major": "Computer Science",
            "semesters": [{"label": "Fall", "courses": [{"code": "CS 201", "title": "Wrong Model Title", "credits": 4}]}],
            "total_credits": 99, "notes": "",
        },
    })
    return client, program


def _doc_text(response):
    document = _exported_document(response)
    return "\n".join(p.text for p in document.paragraphs) + "\n" + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )


@pytest.mark.django_db
@pytest.mark.parametrize("layout", LAYOUTS)
def test_export_uses_program_credits_titles_and_total(django_user_model, layout):
    client, _ = _client_with_program_session(django_user_model)

    response = client.post(EXPORT_URL, {"format": layout}, format="json")

    text = _doc_text(response)
    assert response.status_code == 200
    assert "Data Structures" in text
    assert "Wrong Model Title" not in text
    assert "Total planned credits: 3" in text
    assert "99" not in text


@pytest.mark.django_db
def test_export_rejects_a_plan_from_a_different_program(django_user_model):
    client, program = _client_with_program_session(django_user_model, "export-mismatch-faculty")

    mismatch = client.post(EXPORT_URL, {"program_requirement_id": str(uuid.uuid4())}, format="json")
    match = client.post(EXPORT_URL, {"program_requirement_id": str(program.requirement_id)}, format="json")
    absent = client.post(EXPORT_URL)

    assert mismatch.status_code == 409
    assert "different" in mismatch.data["detail"]
    assert match.status_code == 200
    assert absent.status_code == 200
